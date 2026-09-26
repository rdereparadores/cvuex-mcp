"""Text of the downloaded materials, in parts that can be cited.

Each document is split into numbered parts: pages of a PDF, slides of a
presentation, and headed sections of Word documents, Moodle pages and text
files. There is no OCR: a scanned PDF has no text to extract.
"""

import logging
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path

from cvuex_mcp.formatting import html_to_text, tidy_text

logging.getLogger("pypdf").setLevel(logging.ERROR)  # it warns about every odd PDF

EXTRACTION_VERSION = 1
"""Bump it when extraction changes, so the next sync reads every document again."""
MIN_PAGE_CHARS = 30
"""A page with less text has none worth the name: at most a page number or a header."""
MIN_TEXT_PAGES_SHARE = 0.2
"""A PDF with fewer pages with text than this is taken as scanned."""
SECTION_CHARS = 3000
"""Long parts are split so a search result points close to the match."""


@dataclass(frozen=True)
class Part:
    number: int
    text: str
    heading: str | None = None


@dataclass
class Extraction:
    status: str
    """"ok", "sin_texto" (e.g. a scanned PDF), "formato_antiguo" (.doc, .ppt),
    "no_soportado" or "error"."""
    unit: str = "apartado"
    """What ``Part.number`` counts: "página", "diapositiva" or "apartado"."""
    parts: list[Part] = field(default_factory=list)
    error: str | None = None


# Binary Office formats from before 2007: study material, but unreadable without Office.
LEGACY_FORMATS = {".doc", ".ppt"}
# Why a document can't be searched, for the student.
NOT_SEARCHABLE = {
    "sin_texto": "no tiene texto (quizá es un PDF escaneado)",
    "formato_antiguo": "está en un formato antiguo de Office (.doc o .ppt)",
}


def extract(path: Path) -> Extraction:
    if path.suffix.lower() in LEGACY_FORMATS:
        return Extraction(status="formato_antiguo")
    extractor = EXTRACTORS.get(path.suffix.lower())
    if extractor is None:
        return Extraction(status="no_soportado")
    try:
        return extractor(path)
    except Exception as error:  # a broken file must not stop the rest
        return Extraction(status="error", error=f"{type(error).__name__}: {error}")


def _pdf(path: Path) -> Extraction:
    from pypdf import PdfReader  # imported on demand: it is slow to load

    reader = PdfReader(path)
    if reader.is_encrypted:
        reader.decrypt("")  # teachers often protect them only against editing
    pages = [tidy_text(page.extract_text() or "") for page in reader.pages]
    text_pages = sum(len(text) >= MIN_PAGE_CHARS for text in pages)
    if not text_pages or text_pages < MIN_TEXT_PAGES_SHARE * len(pages):
        return Extraction(status="sin_texto", unit="página")
    parts = [Part(number, text) for number, text in enumerate(pages, start=1) if text]
    return Extraction(status="ok", unit="página", parts=parts)


def _pptx(path: Path) -> Extraction:
    from pptx import Presentation

    parts = []
    for number, slide in enumerate(Presentation(path).slides, start=1):
        texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text
            texts.append(f"Notas: {notes}" if notes.strip() else "")
        title = slide.shapes.title.text if slide.shapes.title is not None else None
        text = tidy_text("\n".join(texts))
        if text:
            parts.append(Part(number, text, tidy_text(title or "") or None))
    return _result("diapositiva", parts)


def _docx(path: Path) -> Extraction:
    from docx import Document

    sections: list[tuple[str | None, list[str]]] = [(None, [])]
    for paragraph in Document(path).paragraphs:
        style = (paragraph.style.name if paragraph.style is not None else "").casefold()
        if style.startswith(("heading", "título", "titulo", "title")) and paragraph.text.strip():
            sections.append((paragraph.text.strip(), []))
        else:
            sections[-1][1].append(paragraph.text)
    return _result("apartado", _numbered(sections))


def _html(path: Path) -> Extraction:
    html = path.read_text(encoding="utf-8", errors="replace")
    pieces = re.split(r"<h[1-4][^>]*>(.*?)</h[1-4]>", html, flags=re.IGNORECASE | re.DOTALL)
    # [before the first heading, heading 1, text 1, heading 2, text 2...]
    sections = [(None, [html_to_text(pieces[0])])]
    for heading, body in zip(pieces[1::2], pieces[2::2], strict=True):
        sections.append((html_to_text(heading) or None, [html_to_text(body)]))
    return _result("apartado", _numbered(sections))


def _text(path: Path) -> Extraction:
    text = path.read_text(encoding="utf-8", errors="replace")
    sections: list[tuple[str | None, list[str]]] = [(None, [])]
    for line in text.splitlines():
        if path.suffix.lower() == ".md" and line.startswith("#"):
            sections.append((line.lstrip("#").strip() or None, []))
        else:
            sections[-1][1].append(line)
    return _result("apartado", _numbered(sections))


def _numbered(sections: list[tuple[str | None, list[str]]]) -> list[Part]:
    """Headed sections → numbered parts, splitting the long ones."""
    parts = []
    for heading, lines in sections:
        text = tidy_text("\n".join(lines))
        for piece in _split(text) if text else ([""] if heading else []):
            parts.append(Part(len(parts) + 1, piece, heading))
    return parts


def _split(text: str) -> list[str]:
    """Pieces of up to ``SECTION_CHARS``, cut between paragraphs when possible."""
    pieces: list[str] = []
    current = ""
    for paragraph in text.split("\n"):
        if current and len(current) + len(paragraph) > SECTION_CHARS:
            pieces.append(current)
            current = ""
        current = f"{current}\n{paragraph}" if current else paragraph
        while len(current) > SECTION_CHARS:  # a single huge paragraph
            pieces.append(current[:SECTION_CHARS])
            current = current[SECTION_CHARS:]
    return [*pieces, current] if current else pieces


def _result(unit: str, parts: list[Part]) -> Extraction:
    has_text = any(part.text for part in parts)
    return Extraction(status="ok" if has_text else "sin_texto", unit=unit, parts=parts)


EXTRACTORS: dict[str, Callable[[Path], Extraction]] = {
    ".pdf": _pdf,
    ".pptx": _pptx,
    ".docx": _docx,
    ".html": _html,
    ".htm": _html,
    ".txt": _text,
    ".md": _text,
}
