"""Text extraction, the search index and reading the materials."""

from pathlib import Path

import pytest
from docx import Document as WordDocument
from helpers import make_pdf
from pptx import Presentation
from pptx.util import Inches

from cvuex_mcp.fulltext import search_words
from cvuex_mcp.materials import extract as extract_module
from cvuex_mcp.materials import search as search_module
from cvuex_mcp.materials.extract import Extraction, Part, extract
from cvuex_mcp.materials.index import DocumentInfo, MaterialsIndex
from cvuex_mcp.materials.search import MaterialNotReadableError, read_material, search_materials

REDES = "Arquitectura de Redes"
URL = "https://campusvirtual.unex.es/zonauex/avuex/mod/resource/view.php?id=1"


def write_docx(path: Path) -> Path:
    document = WordDocument()
    document.add_paragraph("Enunciado de la práctica.")
    document.add_heading("Objetivos", level=1)
    document.add_paragraph("Configurar un servidor DNS.")
    document.add_heading("Entrega", level=2)
    document.add_paragraph("Subir la memoria en PDF.")
    document.save(path)
    return path


def write_pptx(path: Path) -> Path:
    presentation = Presentation()
    for title, body, notes in (
        ("Tema 3", "Encaminamiento estático y dinámico", ""),
        ("OSPF", "Protocolo de estado de enlace", "Preguntan esto en el examen"),
    ):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = title
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
        box.text_frame.text = body
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    presentation.save(path)
    return path


# --- Extraction ---------------------------------------------------------------


def test_pdf_pages(tmp_path):
    path = tmp_path / "tema1.pdf"
    path.write_bytes(make_pdf(["El protocolo TCP garantiza la entrega.", "", "Capa de aplicación"]))
    result = extract(path)
    assert (result.status, result.unit) == ("ok", "página")
    assert result.parts == [
        Part(1, "El protocolo TCP garantiza la entrega."),
        Part(3, "Capa de aplicación"),  # page 2 has no text
    ]


def test_scanned_pdf_has_no_text(tmp_path):
    path = tmp_path / "escaneado.pdf"
    path.write_bytes(make_pdf(["", "", "", "", "12"]))  # at most a page number
    assert extract(path).status == "sin_texto"


def test_word_sections_by_heading(tmp_path):
    result = extract(write_docx(tmp_path / "practica.docx"))
    assert result.unit == "apartado"
    assert result.parts == [
        Part(1, "Enunciado de la práctica."),
        Part(2, "Configurar un servidor DNS.", "Objetivos"),
        Part(3, "Subir la memoria en PDF.", "Entrega"),
    ]


def test_slides_with_titles_and_notes(tmp_path):
    result = extract(write_pptx(tmp_path / "tema3.pptx"))
    assert result.unit == "diapositiva"
    first, second = result.parts
    assert (first.number, first.heading) == (1, "Tema 3")
    assert "Encaminamiento estático" in first.text
    assert second.heading == "OSPF"
    assert "Notas: Preguntan esto en el examen" in second.text


def test_moodle_pages_by_heading(tmp_path):
    path = tmp_path / "guia.html"
    path.write_text(
        "<p>Bienvenida</p><h2>Evaluación</h2><p>Examen: <b>60 %</b></p>"
        "<h3>Prácticas</h3><p>40 %</p>",
        encoding="utf-8",
    )
    assert extract(path).parts == [
        Part(1, "Bienvenida"),
        Part(2, "Examen: 60 %", "Evaluación"),
        Part(3, "40 %", "Prácticas"),
    ]


def test_markdown_and_text(tmp_path):
    (tmp_path / "notas.md").write_text(
        "Intro\n# Uno\nTexto uno\n## Dos\nTexto dos", encoding="utf-8"
    )
    (tmp_path / "leeme.txt").write_text("Solo texto", encoding="utf-8")
    assert [(p.heading, p.text) for p in extract(tmp_path / "notas.md").parts] == [
        (None, "Intro"),
        ("Uno", "Texto uno"),
        ("Dos", "Texto dos"),
    ]
    assert extract(tmp_path / "leeme.txt").parts == [Part(1, "Solo texto")]


def test_long_sections_are_split(tmp_path, monkeypatch):
    monkeypatch.setattr(extract_module, "SECTION_CHARS", 20)
    (tmp_path / "largo.txt").write_text(
        "primer párrafo\nsegundo párrafo\ntercero", encoding="utf-8"
    )
    parts = extract(tmp_path / "largo.txt").parts
    assert [p.text for p in parts] == ["primer párrafo", "segundo párrafo", "tercero"]
    assert [p.number for p in parts] == [1, 2, 3]


def test_other_formats_and_broken_files(tmp_path):
    (tmp_path / "tema1.ppt").write_bytes(b"\xd0\xcf\x11\xe0")
    assert extract(tmp_path / "tema1.ppt").status == "formato_antiguo"
    (tmp_path / "codigo.zip").write_bytes(b"PK")
    (tmp_path / "roto.pdf").write_bytes(b"no es un PDF")
    assert extract(tmp_path / "codigo.zip").status == "no_soportado"
    broken = extract(tmp_path / "roto.pdf")
    assert broken.status == "error"
    assert broken.error


# --- Index and search ---------------------------------------------------------


def info(path: str, course_id: int = 26685, course: str = REDES, stamp: str = "1") -> DocumentInfo:
    return DocumentInfo(
        path=path,
        course_id=course_id,
        course=course,
        section="Tema 3",
        module="Apuntes",
        module_url=URL,
        stamp=stamp,
    )


@pytest.fixture
def index(tmp_path) -> MaterialsIndex:
    index = MaterialsIndex(tmp_path / "materiales.sqlite")
    index.store(
        info(f"{REDES}/03 - Tema 3/tema3.pdf"),
        Extraction(
            "ok",
            "página",
            [
                Part(1, "Introducción al encaminamiento."),
                Part(4, "OSPF es un protocolo de estado de enlace."),
                Part(7, "Los protocolos de vector distancia, como RIP, cuentan saltos."),
            ],
        ),
    )
    index.store(
        info("Bases de Datos/01 - Tema 1/sql.pdf", course_id=15023, course="Bases de Datos"),
        Extraction("ok", "página", [Part(2, "Una transacción cumple las propiedades ACID.")]),
    )
    index.store(info(f"{REDES}/escaneado.pdf"), Extraction("sin_texto", "página"))
    yield index
    index.close()


def test_search_finds_the_page(index, tmp_path):
    result = search_materials(index, tmp_path, "estado de enlace", None, 10)
    [match] = result.resultados
    assert match.model_dump() == {
        "documento_id": 1,
        "documento": "tema3.pdf",
        "asignatura": REDES,
        "seccion": "Tema 3",
        "ubicacion": "página 4",
        "numero": 4,
        "fragmento": "OSPF es un protocolo de «estado» de «enlace».",
        "archivo": str(tmp_path / f"{REDES}/03 - Tema 3/tema3.pdf"),
        "url": URL,
    }
    assert result.todas_las_palabras


def test_search_ignores_accents_case_and_matches_prefixes(index, tmp_path):
    assert search_materials(index, tmp_path, "TRANSACCION", None, 10).resultados
    # "protocolo" also finds "protocolos"
    pages = {m.numero for m in search_materials(index, tmp_path, "protocolo", None, 10).resultados}
    assert pages == {4, 7}


def test_search_falls_back_to_any_word(index, tmp_path):
    result = search_materials(index, tmp_path, "OSPF bicicleta", None, 10)
    assert not result.todas_las_palabras
    assert [m.numero for m in result.resultados] == [4]


def test_search_in_one_course(index, tmp_path):
    assert search_materials(index, tmp_path, "protocolo", 15023, 10).resultados == []


def test_search_warns_about_documents_without_text(index, tmp_path):
    [warning] = search_materials(index, tmp_path, "OSPF", 26685, 10).avisos
    assert warning.startswith("1 documentos no se pueden buscar")


def test_search_without_materials(tmp_path):
    index = MaterialsIndex(tmp_path / "vacio.sqlite")
    result = search_materials(index, tmp_path, "OSPF", None, 10)
    assert (result.resultados, result.avisos) == ([], [search_module.NOTHING_DOWNLOADED])


@pytest.mark.parametrize("query", ['TCP/IP "C++" AND', "NOT OR NEAR(", "*", "de la el"])
def test_any_text_is_a_valid_query(index, tmp_path, query):
    search_materials(index, tmp_path, query, None, 10)  # no FTS5 syntax errors


def test_search_words():
    assert search_words("¿Qué es el protocolo TCP?") == ['"qué"', '"protocolo"*', '"tcp"']


def test_reindexing_keeps_the_document_id(index):
    path = f"{REDES}/03 - Tema 3/tema3.pdf"
    index.store(info(path, stamp="2"), Extraction("ok", "página", [Part(1, "Nueva versión")]))
    assert index.document(1).path == path
    assert index.is_current(path, "2")
    assert index.parts(1, 1, 10) == [(1, None, "Nueva versión")]


# --- Reading ------------------------------------------------------------------


def test_read_pages(index, tmp_path):
    result = read_material(index, tmp_path, 1, 4, None)
    assert result.texto == (
        "[Página 4]\nOSPF es un protocolo de estado de enlace.\n\n"
        "[Página 7]\nLos protocolos de vector distancia, como RIP, cuentan saltos."
    )
    assert (result.desde, result.hasta, result.total, result.continua_en) == (4, 7, 7, None)
    assert (result.unidad, result.documento, result.url) == ("página", "tema3.pdf", URL)


def test_read_continues_where_it_was_cut(index, tmp_path, monkeypatch):
    monkeypatch.setattr(search_module, "MAX_READ_CHARS", 60)
    result = read_material(index, tmp_path, 1, 1, None)
    assert (result.desde, result.hasta, result.continua_en) == (1, 1, 4)


def test_read_a_range(index, tmp_path):
    assert read_material(index, tmp_path, 1, 1, 4).hasta == 4


def test_documents_that_cannot_be_read(index, tmp_path):
    with pytest.raises(MaterialNotReadableError, match="no tiene texto"):
        read_material(index, tmp_path, 3, 1, None)
    with pytest.raises(MaterialNotReadableError, match="No hay ningún documento 99"):
        read_material(index, tmp_path, 99, 1, None)
